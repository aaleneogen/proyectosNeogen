from pathlib import Path
from datetime import datetime
import uuid

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.models import PurchaseOrder
from app.core.config import settings

SAP_COLUMNS = ["SKU", "DESCRIPCION", "CANTIDAD", "PRECIO"]
HEADER_FILL = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
CELL_FONT = Font(name="Calibri", size=10)
BORDER_SIDE = Side(style="thin", color="D1D5DB")
CELL_BORDER = Border(left=BORDER_SIDE, right=BORDER_SIDE, top=BORDER_SIDE, bottom=BORDER_SIDE)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def export_order_to_excel(order: PurchaseOrder) -> str:
    valid_lines = [
        line for line in order.lines
        if line.matched_sku and line.quantity_ordered
    ]

    if not valid_lines:
        raise ValueError("No valid lines to export")

    rows = []
    for line in valid_lines:
        qty = line.quantity_ordered
        rows.append({
            "SKU": line.matched_sku,
            "DESCRIPCION": line.matched_description or "",
            "CANTIDAD": int(qty) if qty == int(qty) else qty,
            "PRECIO": round(line.unit_price, 2) if line.unit_price is not None else 0,
        })

    export_dir = Path(settings.EXPORT_DIR)
    export_dir.mkdir(parents=True, exist_ok=True)

    order_num = (order.order_number or order.id[:8]).replace("/", "-")
    filename = f"SAP_{order_num}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    file_path = export_dir / filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAP Import"

    for col_idx, col_name in enumerate(SAP_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = CELL_BORDER

    for row_idx, row in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=row["SKU"]).alignment = LEFT
        ws.cell(row=row_idx, column=2, value=row["DESCRIPCION"]).alignment = LEFT
        ws.cell(row=row_idx, column=3, value=row["CANTIDAD"]).alignment = CENTER
        price_cell = ws.cell(row=row_idx, column=4, value=row["PRECIO"])
        price_cell.number_format = "#,##0.00"
        price_cell.alignment = CENTER
        for col_idx in range(1, 5):
            ws.cell(row=row_idx, column=col_idx).font = CELL_FONT
            ws.cell(row=row_idx, column=col_idx).border = CELL_BORDER

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 15
    ws.freeze_panes = "A2"

    wb.save(str(file_path))
    return str(file_path)
